import argparse
import json
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import jdatetime
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import (
    DataValidation,
)


BASE_DIR = (
    Path(__file__).resolve().parent
)

OUTPUT_DIR = (
    BASE_DIR
    / "output"
)

STRENGTH_ORDER = {
    "HIGH": 3,
    "MEDIUM": 2,
    "LOW": 1,
    "NONE": 0,
}

PERSIAN_DIGITS = str.maketrans(
    "0123456789",
    "۰۱۲۳۴۵۶۷۸۹",
)

HEADERS = [
    "اولویت",
    "امتیاز",
    "اینستاگرام",
    "تاریخ انتشار",
    "لوکیشن اینستاگرام",
    "شهر تشخیص crawler",
    "استان تشخیص crawler",
    "سیگنال‌های تشخیص",
    "کپشن",
    "لینک پست",
    "وضعیت بررسی",
    "هتل",
    "شهر نهایی",
    "استان نهایی",
    "نام مکان نهایی",
    "عنوان نهایی",
    "نوع مکان",
    "نوع محتوا",
    "خلاصه کپشن",
    "یادداشت",
    "Shortcode",
]

CONTENT_TYPE_OPTIONS = [
    "POST",
    "REEL",
    "STORY",
    "HIGHLIGHT",
    "LIVE",
    "CAROUSEL",
    "IGTV",
    "OTHER",
]

PLACE_TYPE_OPTIONS = [
    "CULTURAL",
    "NATURE",
    "HOTEL",
    "HISTORICAL",
    "RELIGIOUS",
    "URBAN",
    "RURAL",
    "BEACH",
    "MOUNTAIN",
    "DESERT",
    "FOOD",
    "EVENT",
    "OTHER",
]


def resolve_input(
    value: str,
) -> Path:
    path = Path(value)

    if path.exists():
        return path

    candidate = (
        OUTPUT_DIR
        / f"{value}.json"
    )

    if candidate.exists():
        return candidate

    raise FileNotFoundError(
        f"JSON file not found: {value}"
    )


def normalize_signals(
    value,
) -> str:
    if isinstance(
        value,
        list,
    ):
        return "\n".join(
            str(item)
            for item in value
        )

    return str(
        value or ""
    )


def normalize_source_url(
    value,
) -> str:
    return (
        str(value or "")
        .strip()
        .rstrip("/")
    )


def contains_hotel(
    item: dict,
) -> bool:
    if item.get(
        "is_hotel_priority",
        False,
    ):
        return True

    text = " ".join(
        [
            str(
                item.get(
                    "caption",
                    "",
                )
            ),
            str(
                item.get(
                    "instagram_location",
                    "",
                )
            ),
        ]
    ).lower()

    return (
        "هتل" in text
        or "hotel" in text
    )


def get_priority(
    item: dict,
) -> str:
    if contains_hotel(
        item
    ):
        return "HOTEL"

    return item.get(
        "signal_strength",
        "NONE",
    )


def sort_key(
    item: dict,
):
    hotel_priority = (
        1
        if contains_hotel(
            item
        )
        else 0
    )

    strength = item.get(
        "signal_strength",
        "NONE",
    )

    score = item.get(
        "ranking_score",
        item.get(
            "signal_score",
            0,
        ),
    )

    published_at = item.get(
        "published_at",
        "",
    )

    return (
        hotel_priority,
        STRENGTH_ORDER.get(
            strength,
            0,
        ),
        score,
        published_at,
    )


def to_jalali(
    value: str,
) -> str:
    if not value:
        return ""

    try:
        normalized = value.replace(
            "Z",
            "+00:00",
        )

        date_time = (
            datetime.fromisoformat(
                normalized
            )
        )

        if date_time.tzinfo:
            date_time = (
                date_time.astimezone(
                    ZoneInfo(
                        "Asia/Tehran"
                    )
                )
            )

        jalali_date = (
            jdatetime.date.fromgregorian(
                date=date_time.date()
            )
        )

        result = (
            f"{jalali_date.year:04d}/"
            f"{jalali_date.month:02d}/"
            f"{jalali_date.day:02d}"
        )

        return result.translate(
            PERSIAN_DIGITS
        )

    except (
        ValueError,
        TypeError,
    ):
        return value


def guess_hotel(
    item: dict,
) -> str:
    existing = str(
        item.get(
            "hotel",
            "",
        )
    ).strip()

    if existing:
        return existing

    location = str(
        item.get(
            "instagram_location",
            "",
        )
    ).strip()

    if (
        "هتل" in location
        or "hotel" in location.lower()
    ):
        return location

    caption = str(
        item.get(
            "caption",
            "",
        )
    )

    normalized = (
        caption
        .replace("\n", " ")
        .replace("\u200c", " ")
    )

    words = normalized.split()

    for index, word in enumerate(
        words
    ):
        if (
            "هتل" in word
            or word.lower() == "hotel"
        ):
            end = min(
                len(words),
                index + 4,
            )

            return " ".join(
                words[
                    index:end
                ]
            )

    if item.get(
        "is_hotel_priority",
        False,
    ):
        return "نیاز به بررسی"

    return ""


def guess_content_type(
    item: dict,
) -> str:
    existing = str(
        item.get(
            "content_type",
            "",
        )
    ).strip().upper()

    if existing:
        return existing

    product_type = str(
        item.get(
            "product_type",
            "",
        )
    ).strip().lower()

    source_url = normalize_source_url(
        item.get(
            "source_url",
            "",
        )
    ).lower()

    if (
        product_type == "clips"
        or "/reel/" in source_url
    ):
        return "REEL"

    if (
        product_type
        == "carousel_container"
    ):
        return "CAROUSEL"

    if source_url:
        return "POST"

    return ""


def row_values(
    item: dict,
) -> dict:
    return {
        "اولویت": get_priority(
            item
        ),
        "امتیاز": item.get(
            "ranking_score",
            item.get(
                "signal_score",
                0,
            ),
        ),
        "اینستاگرام": item.get(
            "instagram_username",
            "",
        ),
        "لینک پست": item.get(
            "source_url",
            "",
        ),
        "تاریخ انتشار": to_jalali(
            item.get(
                "published_at",
                "",
            )
        ),
        "لوکیشن اینستاگرام": item.get(
            "instagram_location",
            "",
        ),
        "شهر تشخیص crawler": item.get(
            "matched_cities",
            "",
        ),
        "استان تشخیص crawler": item.get(
            "matched_provinces",
            "",
        ),
        "سیگنال‌های تشخیص": normalize_signals(
            item.get(
                "candidate_signals",
                [],
            )
        ),
        "کپشن": item.get(
            "caption",
            "",
        ),
        "وضعیت بررسی": item.get(
            "review_status",
            "pending",
        ),
        "هتل": guess_hotel(
            item
        ),
        "شهر نهایی": (
            item.get(
                "final_cities"
            )
            or item.get(
                "matched_cities",
                "",
            )
        ),
        "استان نهایی": (
            item.get(
                "final_provinces"
            )
            or item.get(
                "matched_provinces",
                "",
            )
        ),
        "عنوان نهایی": item.get(
            "final_title",
            "",
        ),
        "یادداشت": item.get(
            "notes",
            "",
        ),
        "Shortcode": item.get(
            "shortcode",
            "",
        ),
        "نام مکان نهایی": item.get(
            "place_name",
            "",
        ),
        "نوع مکان": item.get(
            "place_type",
            "",
        ),
        "نوع محتوا": guess_content_type(
            item
        ),
        "خلاصه کپشن": item.get(
            "caption_summary",
            "",
        ),
    }


def ensure_column_order(
    workbook,
    sheet,
):
    existing_headers = []

    for column in range(
        1,
        sheet.max_column + 1,
    ):
        value = sheet.cell(
            row=1,
            column=column,
        ).value

        if value is None:
            continue

        existing_headers.append(
            str(value).strip()
        )

    extra_headers = [
        header
        for header
        in existing_headers
        if header not in HEADERS
    ]

    target_headers = (
        HEADERS
        + extra_headers
    )

    existing_rows = []

    if existing_headers:
        header_map = {
            str(
                sheet.cell(
                    row=1,
                    column=column,
                ).value
            ).strip(): column
            for column in range(
                1,
                sheet.max_column + 1,
            )
            if sheet.cell(
                row=1,
                column=column,
            ).value is not None
        }

        for row_number in range(
            2,
            sheet.max_row + 1,
        ):
            existing_rows.append(
                {
                    header: sheet.cell(
                        row=row_number,
                        column=column,
                    ).value
                    for (
                        header,
                        column,
                    )
                    in header_map.items()
                }
            )

    if (
        existing_headers
        == target_headers
    ):
        headers = {
            header: index
            for (
                index,
                header,
            )
            in enumerate(
                target_headers,
                start=1,
            )
        }

        return sheet, headers

    original_title = (
        sheet.title
    )

    original_index = (
        workbook.index(
            sheet
        )
    )

    temporary_title = (
        f"{original_title}"
        "__reordered__"
    )

    if (
        temporary_title
        in workbook.sheetnames
    ):
        workbook.remove(
            workbook[
                temporary_title
            ]
        )

    reordered = (
        workbook.create_sheet(
            temporary_title,
            original_index,
        )
    )

    reordered.append(
        target_headers
    )

    for row in existing_rows:
        reordered.append(
            [
                row.get(
                    header,
                    "",
                )
                for header
                in target_headers
            ]
        )

    workbook.remove(
        sheet
    )

    reordered.title = (
        original_title
    )

    headers = {
        header: index
        for (
            index,
            header,
        )
        in enumerate(
            target_headers,
            start=1,
        )
    }

    return reordered, headers


def existing_source_urls(
    sheet,
    headers: dict[str, int],
) -> set[str]:
    source_column = headers[
        "لینک پست"
    ]

    result = set()

    for row_number in range(
        2,
        sheet.max_row + 1,
    ):
        value = normalize_source_url(
            sheet.cell(
                row=row_number,
                column=source_column,
            ).value
        )

        if value:
            result.add(
                value
            )

    return result


def apply_layout(
    sheet,
    headers: dict[str, int],
):
    sheet.sheet_view.rightToLeft = (
        True
    )

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="1F4E78",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
    )

    thin_side = Side(
        style="thin",
        color="B7B7B7",
    )

    all_borders = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    for column in range(
        1,
        sheet.max_column + 1,
    ):
        cell = sheet.cell(
            row=1,
            column=column,
        )

        cell.fill = header_fill
        cell.font = header_font
        cell.border = all_borders
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    widths = {
        "اولویت": 13,
        "امتیاز": 10,
        "اینستاگرام": 22,
        "لینک پست": 42,
        "تاریخ انتشار": 18,
        "لوکیشن اینستاگرام": 28,
        "شهر تشخیص crawler": 30,
        "استان تشخیص crawler": 35,
        "سیگنال‌های تشخیص": 45,
        "کپشن": 80,
        "وضعیت بررسی": 18,
        "هتل": 30,
        "شهر نهایی": 30,
        "استان نهایی": 30,
        "عنوان نهایی": 40,
        "یادداشت": 35,
        "Shortcode": 20,
        "نام مکان نهایی": 35,
        "نوع مکان": 20,
        "نوع محتوا": 20,
        "خلاصه کپشن": 60,
    }

    for (
        header,
        width,
    ) in widths.items():
        column = headers.get(
            header
        )

        if not column:
            continue

        sheet.column_dimensions[
            get_column_letter(
                column
            )
        ].width = width

    sheet.freeze_panes = "A2"

    if sheet.max_row > 1:
        sheet.auto_filter.ref = (
            f"A1:"
            f"{get_column_letter(sheet.max_column)}"
            f"{sheet.max_row}"
        )


def add_validations(
    sheet,
    headers: dict[str, int],
):
    max_validation_row = max(
        sheet.max_row + 1000,
        5000,
    )

    status_column = get_column_letter(
        headers[
            "وضعیت بررسی"
        ]
    )

    status_validation = DataValidation(
        type="list",
        formula1=(
            '"pending,approved,rejected"'
        ),
        allow_blank=False,
    )

    sheet.add_data_validation(
        status_validation
    )

    status_validation.add(
        f"{status_column}2:"
        f"{status_column}"
        f"{max_validation_row}"
    )

    place_type_column = (
        get_column_letter(
            headers[
                "نوع مکان"
            ]
        )
    )

    place_type_validation = (
        DataValidation(
            type="list",
            formula1=(
                '"'
                + ",".join(
                    PLACE_TYPE_OPTIONS
                )
                + '"'
            ),
            allow_blank=True,
            showErrorMessage=False,
        )
    )

    sheet.add_data_validation(
        place_type_validation
    )

    place_type_validation.add(
        f"{place_type_column}2:"
        f"{place_type_column}"
        f"{max_validation_row}"
    )

    content_type_column = (
        get_column_letter(
            headers[
                "نوع محتوا"
            ]
        )
    )

    content_type_validation = (
        DataValidation(
            type="list",
            formula1=(
                '"'
                + ",".join(
                    CONTENT_TYPE_OPTIONS
                )
                + '"'
            ),
            allow_blank=True,
            showErrorMessage=False,
        )
    )

    sheet.add_data_validation(
        content_type_validation
    )

    content_type_validation.add(
        f"{content_type_column}2:"
        f"{content_type_column}"
        f"{max_validation_row}"
    )


def add_status_formatting(
    sheet,
    headers: dict[str, int],
):
    if sheet.max_row <= 1:
        return

    status_column = (
        get_column_letter(
            headers[
                "وضعیت بررسی"
            ]
        )
    )

    status_range = (
        f"{status_column}2:"
        f"{status_column}"
        f"{sheet.max_row}"
    )

    pending_fill = PatternFill(
        fill_type="solid",
        fgColor="FFF2CC",
    )

    approved_fill = PatternFill(
        fill_type="solid",
        fgColor="C6EFCE",
    )

    rejected_fill = PatternFill(
        fill_type="solid",
        fgColor="FFC7CE",
    )

    sheet.conditional_formatting.add(
        status_range,
        FormulaRule(
            formula=[
                (
                    f'${status_column}2='
                    '"pending"'
                )
            ],
            fill=pending_fill,
        ),
    )

    sheet.conditional_formatting.add(
        status_range,
        FormulaRule(
            formula=[
                (
                    f'${status_column}2='
                    '"approved"'
                )
            ],
            fill=approved_fill,
        ),
    )

    sheet.conditional_formatting.add(
        status_range,
        FormulaRule(
            formula=[
                (
                    f'${status_column}2='
                    '"rejected"'
                )
            ],
            fill=rejected_fill,
        ),
    )


def priority_fills():
    return {
        "HOTEL": PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        ),
        "HIGH": PatternFill(
            fill_type="solid",
            fgColor="E2F0D9",
        ),
        "MEDIUM": PatternFill(
            fill_type="solid",
            fgColor="FFF2CC",
        ),
        "LOW": PatternFill(
            fill_type="solid",
            fgColor="FCE4D6",
        ),
    }


def link_fills():
    return {
        "HOTEL": PatternFill(
            fill_type="solid",
            fgColor="C5D9F1",
        ),
        "HIGH": PatternFill(
            fill_type="solid",
            fgColor="C6E0B4",
        ),
        "MEDIUM": PatternFill(
            fill_type="solid",
            fgColor="FFE699",
        ),
        "LOW": PatternFill(
            fill_type="solid",
            fgColor="F8CBAD",
        ),
    }


def style_data_rows(
    sheet,
    headers: dict[str, int],
):
    fill_by_priority = (
        priority_fills()
    )

    link_fill_by_priority = (
        link_fills()
    )

    thin_side = Side(
        style="thin",
        color="B7B7B7",
    )

    all_borders = Border(
        left=thin_side,
        right=thin_side,
        top=thin_side,
        bottom=thin_side,
    )

    priority_column = (
        headers[
            "اولویت"
        ]
    )

    link_column = (
        headers[
            "لینک پست"
        ]
    )

    for row_number in range(
        2,
        sheet.max_row + 1,
    ):
        priority = str(
            sheet.cell(
                row=row_number,
                column=priority_column,
            ).value
            or ""
        ).strip().upper()

        row_fill = (
            fill_by_priority.get(
                priority
            )
        )

        for column in range(
            1,
            sheet.max_column + 1,
        ):
            cell = sheet.cell(
                row=row_number,
                column=column,
            )

            if row_fill:
                cell.fill = (
                    row_fill
                )

            cell.alignment = (
                Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
            )

            cell.border = (
                all_borders
            )

        link_cell = sheet.cell(
            row=row_number,
            column=link_column,
        )

        source_url = (
            normalize_source_url(
                link_cell.value
            )
        )

        if source_url:
            link_cell.hyperlink = (
                source_url
            )

            link_cell.font = Font(
                color="0563C1",
                underline="single",
            )

        darker_fill = (
            link_fill_by_priority.get(
                priority
            )
        )

        if darker_fill:
            link_cell.fill = (
                darker_fill
            )

        sheet.row_dimensions[
            row_number
        ].height = 90


def append_candidate(
    sheet,
    headers: dict[str, int],
    item: dict,
):
    values = row_values(
        item
    )

    row_number = (
        sheet.max_row + 1
    )

    for (
        header,
        column,
    ) in headers.items():
        if header not in values:
            continue

        sheet.cell(
            row=row_number,
            column=column,
            value=values[
                header
            ],
        )


def json_to_excel(
    input_path: Path,
    output_path: Path,
):
    with input_path.open(
        "r",
        encoding="utf-8",
    ) as file:
        data = json.load(file)

    if not isinstance(
        data,
        list,
    ):
        raise ValueError(
            "JSON root must be a list."
        )

    data.sort(
        key=sort_key,
        reverse=True,
    )

    output_path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    existed = (
        output_path.exists()
    )

    if existed:
        workbook = load_workbook(
            output_path
        )

        if (
            "Candidates"
            in workbook.sheetnames
        ):
            sheet = workbook[
                "Candidates"
            ]
        else:
            sheet = workbook.create_sheet(
                "Candidates"
            )
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Candidates"

    sheet, headers = (
        ensure_column_order(
            workbook,
            sheet,
        )
    )

    known_urls = existing_source_urls(
        sheet,
        headers,
    )

    added = 0
    skipped = 0

    for item in data:
        source_url = (
            normalize_source_url(
                item.get(
                    "source_url",
                    "",
                )
            )
        )

        if (
            source_url
            and source_url
            in known_urls
        ):
            skipped += 1
            continue

        append_candidate(
            sheet,
            headers,
            item,
        )

        added += 1

        if source_url:
            known_urls.add(
                source_url
            )

    apply_layout(
        sheet,
        headers,
    )

    style_data_rows(
        sheet,
        headers,
    )

    add_validations(
        sheet,
        headers,
    )

    add_status_formatting(
        sheet,
        headers,
    )

    workbook.save(
        output_path
    )

    print(
        (
            "Excel updated:"
            if existed
            else "Excel created:"
        )
    )

    print(
        output_path
    )

    print(
        "JSON candidates:",
        len(data),
    )

    print(
        "Added:",
        added,
    )

    print(
        "Skipped existing:",
        skipped,
    )

    print(
        "Workbook rows:",
        max(
            sheet.max_row - 1,
            0,
        ),
    )


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Convert Instagram "
            "candidate JSON "
            "to a merge-safe Excel "
            "review workbook."
        )
    )

    parser.add_argument(
        "input",
        help=(
            "Username or JSON path"
        ),
    )

    parser.add_argument(
        "--output",
        help=(
            "Optional XLSX "
            "output path"
        ),
    )

    args = parser.parse_args()

    input_path = resolve_input(
        args.input
    )

    if args.output:
        output_path = Path(
            args.output
        )
    else:
        output_path = (
            input_path.with_suffix(
                ".xlsx"
            )
        )

    json_to_excel(
        input_path,
        output_path,
    )


if __name__ == "__main__":
    main()
