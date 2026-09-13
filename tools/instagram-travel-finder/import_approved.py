import argparse
import hashlib
import json
import os
import re
import shutil
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

DEFAULT_API_BASE = (
    "http://localhost:4000/api/v1"
)

API_BASE = os.getenv(
    "HOTELYAB_API_BASE",
    DEFAULT_API_BASE,
).rstrip("/")

ADMIN_COOKIE = os.getenv(
    "HOTELYAB_ADMIN_COOKIE",
    "",
)

REVIEW_PLACEHOLDERS = {
    "نیاز به بررسی",
    "needs review",
}

REPO_ROOT = Path(__file__).resolve().parents[2]
PUBLIC_DIR = REPO_ROOT / "apps" / "web" / "public"

CONTENT_TYPE_OPTIONS = {
    "POST",
    "REEL",
    "VIDEO_POST",
    "STORY",
    "HIGHLIGHT",
    "LIVE",
    "CAROUSEL",
    "IGTV",
    "OTHER",
}


def content_kind_for(
    content_type: str,
) -> str:
    normalized = str(
        content_type or ""
    ).strip().upper()

    if normalized in {
        "STORY",
        "HIGHLIGHT",
    }:
        return "STORY"

    if normalized in {
        "POST",
        "CAROUSEL",
    }:
        return "POST"

    return "VIDEO"

PLACE_TYPE_OPTIONS = {
    "CULTURAL",
    "NATURE",
    "HOTEL",
    "HISTORICAL",
    "RELIGIOUS",
    "URBAN",
    "RURAL",
    "BEACH",
    "MOUNTAIN",
    "DESERT",
    "FOOD",
    "EVENT",
    "OTHER",
}

class Color:
    RESET = "\033[0m"
    BOLD = "\033[1m"

    RED = "\033[91m"
    GREEN = "\033[92m"
    YELLOW = "\033[93m"
    BLUE = "\033[94m"
    MAGENTA = "\033[95m"
    CYAN = "\033[96m"
    GRAY = "\033[90m"
    WHITE = "\033[97m"


def color(
    text,
    value,
):
    return (
        f"{value}{text}"
        f"{Color.RESET}"
    )


def label(
    name: str,
    value: str,
):
    print(
        color(
            f"{name:<14}",
            Color.GRAY,
        ),
        value,
    )

def normalize_text(
    value,
) -> str:
    text = str(
        value or ""
    ).strip()

    replacements = {
        "ي": "ی",
        "ى": "ی",
        "ك": "ک",
        "\u200c": " ",
        "\u200f": " ",
        "\u200e": " ",
    }

    for old, new in replacements.items():
        text = text.replace(
            old,
            new,
        )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip().casefold()


def normalize_handle(
    value,
) -> str:
    return (
        normalize_text(value)
        .lstrip("@")
    )


def split_multi(
    value,
) -> list[str]:
    text = str(
        value or ""
    ).strip()

    if not text:
        return []

    return [
        part.strip()
        for part in re.split(
            r"\s*\|\s*",
            text,
        )
        if part.strip()
    ]


def api_get_json(
    path: str,
):
    if not ADMIN_COOKIE:
        raise RuntimeError(
            "HOTELYAB_ADMIN_COOKIE "
            "is not set."
        )

    url = (
        f"{API_BASE}"
        f"{path}"
    )

    request = Request(
        url,
        headers={
            "Accept": (
                "application/json"
            ),
            "Cookie": ADMIN_COOKIE,
        },
    )

    try:
        with urlopen(
            request,
            timeout=30,
        ) as response:
            return json.loads(
                response.read()
                .decode("utf-8")
            )

    except HTTPError as exc:
        body = (
            exc.read()
            .decode(
                "utf-8",
                errors="replace",
            )
        )

        raise RuntimeError(
            "Hotel-Yab API returned "
            f"HTTP {exc.code}: "
            f"{body}"
        ) from exc

    except URLError as exc:
        raise RuntimeError(
            "Could not connect to "
            f"Hotel-Yab API: {exc}"
        ) from exc



def api_post_json(
    path: str,
    payload: dict,
):
    if not ADMIN_COOKIE:
        raise RuntimeError(
            "HOTELYAB_ADMIN_COOKIE "
            "is not set."
        )

    url = (
        f"{API_BASE}"
        f"{path}"
    )

    body = json.dumps(
        payload,
        ensure_ascii=False,
        separators=(
            ",",
            ":",
        ),
    ).encode(
        "utf-8"
    )

    request = Request(
        url,
        data=body,
        method="POST",
        headers={
            "Accept": (
                "application/json"
            ),
            "Content-Type": (
                "application/json"
            ),
            "Cookie": ADMIN_COOKIE,
        },
    )

    try:
        with urlopen(
            request,
            timeout=45,
        ) as response:
            raw = response.read()

            if not raw:
                return {}

            return json.loads(
                raw.decode(
                    "utf-8"
                )
            )

    except HTTPError as exc:
        body_text = (
            exc.read()
            .decode(
                "utf-8",
                errors="replace",
            )
        )

        raise RuntimeError(
            "Hotel-Yab API returned "
            f"HTTP {exc.code}: "
            f"{body_text}"
        ) from exc

    except URLError as exc:
        raise RuntimeError(
            "Could not connect to "
            f"Hotel-Yab API: {exc}"
        ) from exc


def get_collection(
    payload: dict,
    *names: str,
) -> list[dict]:
    root = payload.get(
        "data",
        payload,
    )

    for name in names:
        value = root.get(name)

        if isinstance(
            value,
            list,
        ):
            return value

    return []


def load_catalog():
    payload = api_get_json(
        "/admin/catalog/bootstrap"
    )

    return {
        "destinations": (
            get_collection(
                payload,
                "destinations",
            )
        ),
        "hotels": (
            get_collection(
                payload,
                "hotels",
            )
        ),
        "people": (
            get_collection(
                payload,
                "notablePeople",
                "notable_people",
                "people",
            )
        ),
        "videos": (
            get_collection(
                payload,
                "videos",
            )
        ),
    }


def find_person(
    people: list[dict],
    username: str,
):
    target = normalize_handle(
        username
    )

    for person in people:
        handle = (
            person.get(
                "instagramHandle"
            )
            or person.get(
                "instagram_handle"
            )
            or ""
        )

        if (
            normalize_handle(handle)
            == target
        ):
            return person

    return None


def normalize_hotel_name(
    value,
) -> str:
    text = normalize_text(
        value
    )

    text = re.sub(
        r"[\u0640\u064b-\u065f\u0670\u06d6-\u06ed]",
        "",
        text,
    )

    text = re.sub(
        r"^(?:هتل|hotel)\s+",
        "",
        text,
        flags=re.IGNORECASE,
    )

    text = re.sub(
        r"\s+",
        " ",
        text,
    )

    return text.strip()


def find_hotel(
    hotels: list[dict],
    hotel_name: str,
):
    target = normalize_hotel_name(
        hotel_name
    )

    for hotel in hotels:
        name = hotel.get(
            "name",
            "",
        )

        if (
            normalize_hotel_name(name)
            == target
        ):
            return hotel

    return None


def destination_type(
    destination: dict,
) -> str:
    value = normalize_text(
        destination.get(
            "type",
            "",
        )
    )

    if value in {
        "city",
        "cities",
    }:
        return "CITY"

    if value in {
        "province",
        "provinces",
    }:
        return "PROVINCE"

    return value.upper()


def find_destination(
    destinations: list[dict],
    name: str,
    expected_type: str,
):
    target = normalize_text(
        name
    )

    for destination in destinations:
        if (
            destination_type(
                destination
            )
            != expected_type
        ):
            continue

        destination_name = (
            destination.get(
                "name",
                "",
            )
        )

        if (
            normalize_text(
                destination_name
            )
            == target
        ):
            return destination

    return None


def find_video_by_source(
    videos: list[dict],
    source_url: str,
):
    target = (
        str(source_url or "")
        .strip()
        .rstrip("/")
    )

    for video in videos:
        existing = (
            str(
                video.get(
                    "sourceUrl",
                    "",
                )
            )
            .strip()
            .rstrip("/")
        )

        if (
            existing
            and existing == target
        ):
            return video

    return None


def existing_video_title_keys(
    videos: list[dict],
) -> set[str]:
    return {
        normalize_text(video.get("title", ""))
        for video in videos
        if normalize_text(video.get("title", ""))
    }


def allocate_unique_video_title(
    title: str,
    used_titles: set[str],
) -> tuple[str, bool]:
    base = str(title or "").strip()
    candidate = base
    suffix = 2

    while normalize_text(candidate) in used_titles:
        suffix_text = f" {suffix}"
        candidate = f"{base[: 240 - len(suffix_text)].rstrip()}{suffix_text}"
        suffix += 1

    used_titles.add(normalize_text(candidate))
    return candidate, candidate != base


def load_excel_rows(
    path: Path,
):
    if path.suffix.lower() == ".json":
        return load_review_export_rows(
            path
        )

    workbook = load_workbook(
        path,
        data_only=True,
    )

    sheet = (
        workbook[
            "Candidates"
        ]
        if "Candidates"
        in workbook.sheetnames
        else workbook.active
    )

    headers = {}

    for column in range(
        1,
        sheet.max_column + 1,
    ):
        value = sheet.cell(
            row=1,
            column=column,
        ).value

        if value is not None:
            headers[
                str(value).strip()
            ] = column

    required_headers = [
        "اینستاگرام",
        "لینک پست",
        "تاریخ انتشار",
        "وضعیت بررسی",
        "هتل",
        "شهر نهایی",
        "استان نهایی",
        "نام مکان نهایی",
        "عنوان نهایی",
        "نوع مکان",
        "نوع محتوا",
        "خلاصه کپشن",
        "یادداشت",
        "Shortcode",
    ]

    missing = [
        header
        for header
        in required_headers
        if header not in headers
    ]

    if missing:
        raise RuntimeError(
            "Missing Excel columns: "
            + ", ".join(missing)
        )

    rows = []

    for row_number in range(
        2,
        sheet.max_row + 1,
    ):
        row = {
            header: sheet.cell(
                row=row_number,
                column=column,
            ).value
            for (
                header,
                column,
            )
            in headers.items()
        }

        row["_row_number"] = (
            row_number
        )

        rows.append(row)

    return rows


def load_review_export_rows(
    path: Path,
):
    payload = json.loads(
        path.read_text(
            encoding="utf-8"
        )
    )

    if (
        not isinstance(payload, dict)
        or payload.get("kind")
        != "hotel-yab-crawl-review"
    ):
        raise RuntimeError(
            "JSON input is not a Hotel-Yab "
            "crawl review export."
        )

    rows = payload.get("rows")

    if not isinstance(rows, list):
        raise RuntimeError(
            "Review export has no rows list."
        )

    required_headers = [
        "اینستاگرام",
        "لینک پست",
        "تاریخ انتشار",
        "وضعیت بررسی",
        "هتل",
        "شهر نهایی",
        "استان نهایی",
        "نام مکان نهایی",
        "عنوان نهایی",
        "نوع مکان",
        "نوع محتوا",
        "خلاصه کپشن",
        "یادداشت",
        "Shortcode",
    ]

    normalized = []

    for index, raw_row in enumerate(
        rows,
        start=2,
    ):
        if not isinstance(raw_row, dict):
            raise RuntimeError(
                "Review export row "
                f"{index} is not an object."
            )

        missing = [
            header
            for header in required_headers
            if header not in raw_row
        ]

        if missing:
            raise RuntimeError(
                "Review export row "
                f"{index} is missing: "
                + ", ".join(missing)
            )

        row = dict(raw_row)
        row["_row_number"] = int(
            row.get(
                "_row_number",
                index,
            )
        )
        normalized.append(row)

    return normalized


def is_approved(
    row: dict,
) -> bool:
    return (
        normalize_text(
            row.get(
                "وضعیت بررسی"
            )
        )
        == "approved"
    )


def clean_hotel_name(
    value,
) -> str:
    value = str(
        value or ""
    ).strip()

    if (
        normalize_text(value)
        in {
            normalize_text(item)
            for item
            in REVIEW_PLACEHOLDERS
        }
    ):
        return ""

    return value


def clean_cell(
    row: dict,
    key: str,
) -> str:
    return str(
        row.get(
            key,
            "",
        )
        or ""
    ).strip()


def extract_video_sequence(
    value,
) -> int | None:
    match = re.search(
        r"(?:^|-)(\d+)$",
        str(
            value or ""
        ).strip(),
    )

    if not match:
        return None

    return int(
        match.group(1)
    )


def existing_sequence_max(
    catalog: dict,
    username: str,
    person_slug: str,
) -> int:
    maximum = 0

    for video in catalog[
        "videos"
    ]:
        video_username = (
            video.get(
                "instagramUsername"
            )
            or video.get(
                "instagram_username"
            )
            or ""
        )

        if (
            normalize_handle(
                video_username
            )
            != normalize_handle(
                username
            )
        ):
            continue

        sequence = (
            extract_video_sequence(
                video.get(
                    "id"
                )
            )
        )

        if sequence is not None:
            maximum = max(
                maximum,
                sequence,
            )

    for folder_name in (
        "travel-videos",
        "hotel-videos",
    ):
        folder = (
            PUBLIC_DIR
            / folder_name
            / person_slug
        )

        if not folder.exists():
            continue

        for path in folder.iterdir():
            if not path.is_file():
                continue

            if folder_name == "travel-videos":
                match = re.match(
                    r"^(\d{3})(?:-\d{2})?(?:-thumbnail)?\.",
                    path.name,
                )
            else:
                match = re.search(
                    r"-(\d{3})(?:-\d{2})?(?:-thumbnail)?\.",
                    path.name,
                )

            if match:
                maximum = max(
                    maximum,
                    int(
                        match.group(1)
                    ),
                )

    return maximum


def allocate_video_identity(
    catalog: dict,
    username: str,
    person_slug: str,
    hotel_record: dict | None,
    sequence_state: dict[str, int],
) -> dict:
    key = normalize_handle(
        username
    )

    if key not in sequence_state:
        sequence_state[
            key
        ] = existing_sequence_max(
            catalog,
            username,
            person_slug,
        )

    sequence_state[
        key
    ] += 1

    sequence = sequence_state[
        key
    ]

    sequence_text = (
        f"{sequence:03d}"
    )

    video_id = (
        f"{username}-"
        f"{sequence_text}"
    )

    if hotel_record:
        hotel_slug = str(
            hotel_record.get(
                "slug",
                "",
            )
            or ""
        ).strip()

        stem = (
            f"/hotel-videos/"
            f"{person_slug}/"
            f"{hotel_slug}-"
            f"{sequence_text}"
        )
    else:
        stem = (
            f"/travel-videos/"
            f"{person_slug}/"
            f"{sequence_text}"
        )

    return {
        "sequence": sequence_text,
        "video_id": video_id,
        "stem": stem,
    }


def unique_ids(
    records: list[dict],
) -> list[str]:
    result = []
    seen = set()

    for record in records:
        value = str(
            record.get(
                "id",
                "",
            )
            or ""
        ).strip()

        if (
            not value
            or value in seen
        ):
            continue

        seen.add(
            value
        )
        result.append(
            value
        )

    return result


def analyze_row(
    row: dict,
    catalog: dict,
    sequence_state: dict[str, int],
    batch_sources: set[str],
    used_titles: set[str],
):
    problems = []
    warnings = []

    row_number = row[
        "_row_number"
    ]

    username = clean_cell(
        row,
        "اینستاگرام",
    )

    source_url = clean_cell(
        row,
        "لینک پست",
    )

    normalized_source = (
        source_url.rstrip("/")
    )

    hotel_name = (
        clean_hotel_name(
            row.get(
                "هتل"
            )
        )
    )

    video_category = (
        "HOTEL"
        if hotel_name
        else "TRAVEL"
    )

    city_names = split_multi(
        row.get(
            "شهر نهایی"
        )
    )

    province_names = split_multi(
        row.get(
            "استان نهایی"
        )
    )

    title = clean_cell(
        row,
        "عنوان نهایی",
    )

    place_name = clean_cell(
        row,
        "نام مکان نهایی",
    )

    place_type = clean_cell(
        row,
        "نوع مکان",
    ).upper()

    content_type = clean_cell(
        row,
        "نوع محتوا",
    ).upper()

    content_kind = content_kind_for(
        content_type
    )

    published_date = clean_cell(
        row,
        "تاریخ انتشار",
    )

    caption_summary = clean_cell(
        row,
        "خلاصه کپشن",
    )

    notes = clean_cell(
        row,
        "یادداشت",
    )

    if not username:
        problems.append(
            "Instagram username "
            "is empty"
        )

    if not source_url:
        problems.append(
            "source_url is empty"
        )
    elif not source_url.startswith(
        (
            "http://",
            "https://",
        )
    ):
        problems.append(
            "source_url must be "
            "an HTTP(S) URL"
        )

    if (
        normalized_source
        and normalized_source
        in batch_sources
    ):
        problems.append(
            "Duplicate source_url "
            "inside approved rows"
        )

    if normalized_source:
        batch_sources.add(
            normalized_source
        )

    existing_video = None

    if source_url:
        existing_video = (
            find_video_by_source(
                catalog[
                    "videos"
                ],
                source_url,
            )
        )

    if existing_video:
        return {
            "row": row_number,
            "status": (
                "ALREADY_EXISTS"
            ),
            "username": username,
            "source_url": source_url,
            "hotel": hotel_name,
            "video_category": (
                video_category
            ),
            "cities": city_names,
            "provinces": (
                province_names
            ),
            "person": None,
            "creator_slug": "",
            "destinations": [],
            "hotel_record": None,
            "payload": None,
            "problems": [],
            "warnings": [],
        }

    person = None
    creator_slug = ""

    if username:
        person = find_person(
            catalog[
                "people"
            ],
            username,
        )

        if not person:
            problems.append(
                "Missing person for "
                f"@{username}"
            )
        else:
            creator_slug = str(
                person.get(
                    "slug",
                    "",
                )
                or ""
            ).strip()

            if not creator_slug:
                problems.append(
                    "Person slug is empty"
                )

    destination_records = []

    for city_name in city_names:
        destination = (
            find_destination(
                catalog[
                    "destinations"
                ],
                city_name,
                "CITY",
            )
        )

        if destination:
            destination_records.append(
                destination
            )
        else:
            problems.append(
                "Missing city: "
                f"{city_name}"
            )

    for province_name in (
        province_names
    ):
        destination = (
            find_destination(
                catalog[
                    "destinations"
                ],
                province_name,
                "PROVINCE",
            )
        )

        if destination:
            destination_records.append(
                destination
            )
        else:
            problems.append(
                "Missing province: "
                f"{province_name}"
            )

    if not (
        city_names
        or province_names
    ):
        problems.append(
            "No final city "
            "or province"
        )

    hotel_record = None

    if hotel_name:
        hotel_record = find_hotel(
            catalog[
                "hotels"
            ],
            hotel_name,
        )

        if not hotel_record:
            problems.append(
                "Missing hotel: "
                f"{hotel_name}"
            )

    if not title:
        problems.append(
            "Final title is empty"
        )
    elif len(
        title
    ) > 240:
        problems.append(
            "Final title is longer "
            "than 240 characters"
        )

    if not place_name:
        problems.append(
            "Final place name "
            "is empty"
        )
    elif len(
        place_name
    ) > 200:
        problems.append(
            "Final place name is "
            "longer than 200 characters"
        )

    if not place_type:
        problems.append(
            "Place type is empty"
        )
    elif place_type not in (
        PLACE_TYPE_OPTIONS
    ):
        problems.append(
            "Unsupported place type: "
            f"{place_type}"
        )

    if not content_type:
        problems.append(
            "Content type is empty"
        )
    elif content_type not in (
        CONTENT_TYPE_OPTIONS
    ):
        problems.append(
            "Unsupported content type: "
            f"{content_type}"
        )

    if not caption_summary:
        problems.append(
            "Caption summary is empty"
        )

    if (
        published_date
        and len(
            published_date
        ) > 20
    ):
        problems.append(
            "Published date is longer "
            "than 20 characters"
        )

    if not published_date:
        warnings.append(
            "Published date is empty"
        )

    destination_ids = unique_ids(
        destination_records
    )

    hotel_ids = (
        unique_ids(
            [
                hotel_record
            ]
        )
        if hotel_record
        else []
    )

    if (
        video_category
        == "TRAVEL"
        and not destination_ids
    ):
        problems.append(
            "TRAVEL video needs at "
            "least one destination"
        )

    if (
        video_category
        == "HOTEL"
        and not hotel_ids
    ):
        problems.append(
            "HOTEL video needs "
            "a valid hotel"
        )

    identity = None
    payload = None

    if not problems:
        title, title_was_suffixed = allocate_unique_video_title(
            title,
            used_titles,
        )

        if title_was_suffixed:
            warnings.append(
                "Duplicate title was renamed to: "
                f"{title}"
            )

        identity = (
            allocate_video_identity(
                catalog,
                username,
                creator_slug,
                hotel_record,
                sequence_state,
            )
        )

        payload = {
            "id": identity[
                "video_id"
            ],
            "videoCategory": (
                video_category
            ),
            "contentKind": (
                content_kind
            ),
            "instagramUsername": (
                username
            ),
            "platform": "INSTAGRAM",
            "personCategory": None,
            "contentType": (
                content_type
            ),
            "sourceUrl": (
                source_url
            ),
            "title": title,
            "placeName": (
                place_name
            ),
            "placeType": (
                place_type
            ),
            "publishedDate": (
                published_date
                or None
            ),
            "captionSummary": (
                caption_summary
                or None
            ),
            "evidenceType": (
                "ORIGINAL_POST"
            ),
            "verificationStatus": (
                "VERIFIED"
            ),
            "notes": (
                notes
                or None
            ),
            "mediaUrl": "",
            "thumbnailUrl": "",
            "publicationStatus": (
                "PUBLISHED"
            ),
            "destinationIds": (
                destination_ids
            ),
            "hotelIds": hotel_ids,
        }

    status = (
        "BLOCKED"
        if problems
        else "READY"
    )

    return {
        "row": row_number,
        "status": status,
        "username": username,
        "source_url": source_url,
        "hotel": hotel_name,
        "video_category": (
            video_category
        ),
        "cities": city_names,
        "provinces": province_names,
        "person": person,
        "creator_slug": (
            creator_slug
        ),
        "destinations": (
            destination_records
        ),
        "hotel_record": (
            hotel_record
        ),
        "title": title,
        "place_name": place_name,
        "place_type": place_type,
        "content_type": content_type,
        "content_kind": content_kind,
        "identity": identity,
        "published_date": (
            published_date
        ),
        "caption_summary": (
            caption_summary
        ),
        "payload": payload,
        "problems": problems,
        "warnings": warnings,
    }


def print_result(
    result: dict,
):
    status = result[
        "status"
    ]

    row_number = result[
        "row"
    ]

    print()

    if status == "READY":
        status_text = color(
            (
                f"✅ ROW {row_number} "
                "— READY"
            ),
            Color.GREEN
            + Color.BOLD,
        )

    elif status == "BLOCKED":
        status_text = color(
            (
                f"❌ ROW {row_number} "
                "— BLOCKED"
            ),
            Color.RED
            + Color.BOLD,
        )

    elif (
        status
        == "ALREADY_EXISTS"
    ):
        status_text = color(
            (
                f"⏭ ROW {row_number} "
                "— ALREADY EXISTS"
            ),
            Color.CYAN
            + Color.BOLD,
        )

    else:
        status_text = (
            f"ROW {row_number} "
            f"— {status}"
        )

    separator = (
        "━━━━━━━━━━━━━━━━━━━━━━━━"
        "━━━━━━━━━━━━━━━━━━━━"
    )

    print(
        color(
            separator,
            Color.GRAY,
        )
    )

    print(
        status_text
    )

    print(
        color(
            separator,
            Color.GRAY,
        )
    )

    label(
        "Instagram",
        (
            result[
                "username"
            ]
            or "—"
        ),
    )

    label(
        "Source",
        color(
            (
                result[
                    "source_url"
                ]
                or "—"
            ),
            Color.CYAN,
        ),
    )

    if (
        status
        == "ALREADY_EXISTS"
    ):
        print()

        print(
            color(
                (
                    "This video already "
                    "exists in Hotel-Yab."
                ),
                Color.CYAN,
            )
        )

        print(
            color(
                (
                    "Nothing will be "
                    "created for this row."
                ),
                Color.GRAY,
            )
        )

        return

    label(
        "Category",
        color(
            result[
                "video_category"
            ],
            (
                Color.BLUE
                if result[
                    "video_category"
                ]
                == "HOTEL"
                else Color.MAGENTA
            ),
        ),
    )

    person = result[
        "person"
    ]

    if person:
        label(
            "Person",
            color(
                "✅ FOUND",
                Color.GREEN,
            ),
        )

        label(
            "Creator slug",
            result[
                "creator_slug"
            ]
            or "—",
        )

    else:
        label(
            "Person",
            color(
                "❌ NOT FOUND",
                Color.RED,
            ),
        )

    hotel_name = result[
        "hotel"
    ]

    if hotel_name:
        if result[
            "hotel_record"
        ]:
            label(
                "Hotel",
                color(
                    (
                        f"✅ "
                        f"{hotel_name} "
                        "— FOUND"
                    ),
                    Color.GREEN,
                ),
            )

        else:
            label(
                "Hotel",
                color(
                    (
                        f"❌ "
                        f"{hotel_name} "
                        "— NOT FOUND"
                    ),
                    Color.RED,
                ),
            )

    else:
        label(
            "Hotel",
            color(
                "—",
                Color.GRAY,
            ),
        )

    if result[
        "cities"
    ]:
        city_values = []

        for city_name in (
            result[
                "cities"
            ]
        ):
            found = (
                find_destination(
                    catalog_for_print[
                        "destinations"
                    ],
                    city_name,
                    "CITY",
                )
            )

            city_values.append(
                color(
                    (
                        f"✅ {city_name}"
                        if found
                        else (
                            f"❌ "
                            f"{city_name}"
                        )
                    ),
                    (
                        Color.GREEN
                        if found
                        else Color.RED
                    ),
                )
            )

        label(
            "City",
            " | ".join(
                city_values
            ),
        )

    else:
        label(
            "City",
            color(
                "—",
                Color.GRAY,
            ),
        )

    if result[
        "provinces"
    ]:
        province_values = []

        for province_name in (
            result[
                "provinces"
            ]
        ):
            found = (
                find_destination(
                    catalog_for_print[
                        "destinations"
                    ],
                    province_name,
                    "PROVINCE",
                )
            )

            province_values.append(
                color(
                    (
                        "✅ "
                        f"{province_name}"
                        if found
                        else (
                            "❌ "
                            f"{province_name}"
                        )
                    ),
                    (
                        Color.GREEN
                        if found
                        else Color.RED
                    ),
                )
            )

        label(
            "Province",
            " | ".join(
                province_values
            ),
        )

    else:
        label(
            "Province",
            color(
                "—",
                Color.GRAY,
            ),
        )

    label(
        "Title",
        result.get(
            "title"
        )
        or "—",
    )

    label(
        "Place",
        result.get(
            "place_name"
        )
        or "—",
    )

    label(
        "Place type",
        result.get(
            "place_type"
        )
        or "—",
    )

    label(
        "Content type",
        result.get(
            "content_type"
        )
        or "—",
    )

    label(
        "Published",
        result.get(
            "published_date"
        )
        or "—",
    )

    label(
        "Caption",
        result.get(
            "caption_summary"
        )
        or "—",
    )

    payload = result.get(
        "payload"
    )

    if payload:
        print()

        print(
            color(
                "FINAL PAYLOAD:",
                Color.WHITE
                + Color.BOLD,
            )
        )

        label(
            "Video ID",
            payload[
                "id"
            ],
        )

        label(
            "Media",
            color(
                payload[
                    "mediaUrl"
                ],
                (
                    Color.BLUE
                    if result[
                        "video_category"
                    ]
                    == "HOTEL"
                    else Color.MAGENTA
                ),
            ),
        )

        label(
            "Thumbnail",
            payload[
                "thumbnailUrl"
            ],
        )

        label(
            "Dest IDs",
            (
                " | ".join(
                    payload[
                        "destinationIds"
                    ]
                )
                or "—"
            ),
        )

        label(
            "Hotel IDs",
            (
                " | ".join(
                    payload[
                        "hotelIds"
                    ]
                )
                or "—"
            ),
        )

        label(
            "Verify",
            payload[
                "verificationStatus"
            ],
        )

        label(
            "Publish",
            payload[
                "publicationStatus"
            ],
        )

    if result[
        "warnings"
    ]:
        print()

        print(
            color(
                "WARNINGS:",
                Color.YELLOW
                + Color.BOLD,
            )
        )

        for warning in result[
            "warnings"
        ]:
            print(
                color(
                    f"  ⚠ {warning}",
                    Color.YELLOW,
                )
            )

    if result[
        "problems"
    ]:
        print()

        print(
            color(
                "BLOCK REASONS:",
                Color.RED
                + Color.BOLD,
            )
        )

        for problem in result[
            "problems"
        ]:
            print(
                color(
                    f"  ❌ {problem}",
                    Color.RED,
                )
            )



def resolve_excel_path(
    value: str,
) -> Path:
    direct = Path(
        value
    ).expanduser()

    if direct.exists():
        return direct.resolve()

    output_dir = (
        Path(__file__)
        .resolve()
        .parent
        / "output"
    )

    filenames = (
        [value]
        if value.lower().endswith(
            (".xlsx", ".json")
        )
        else [
            f"{value}.xlsx",
            f"{value}.reviewed.json",
        ]
    )

    for filename in filenames:
        candidate = (
            output_dir
            / filename
        )

        if candidate.exists():
            return candidate.resolve()

    raise FileNotFoundError(
        direct
    )


def analyze_approved_rows(
    excel_path: Path,
    catalog: dict,
):
    rows = load_excel_rows(
        excel_path
    )

    approved = [
        row
        for row in rows
        if is_approved(row)
    ]

    sequence_state = {}
    batch_sources = set()
    used_titles = existing_video_title_keys(catalog["videos"])

    results = [
        analyze_row(
            row,
            catalog,
            sequence_state,
            batch_sources,
            used_titles,
        )
        for row in approved
    ]

    return approved, results


def approved_rows_fingerprint(
    approved: list[dict],
) -> str:
    keys = [
        "اینستاگرام",
        "لینک پست",
        "تاریخ انتشار",
        "هتل",
        "شهر نهایی",
        "استان نهایی",
        "نام مکان نهایی",
        "عنوان نهایی",
        "نوع مکان",
        "نوع محتوا",
        "خلاصه کپشن",
        "یادداشت",
        "Shortcode",
    ]

    payload = [
        {
            "row": row[
                "_row_number"
            ],
            **{
                key: clean_cell(
                    row,
                    key,
                )
                for key in keys
            },
        }
        for row in approved
    ]

    encoded = json.dumps(
        payload,
        ensure_ascii=False,
        sort_keys=True,
        separators=(
            ",",
            ":",
        ),
    ).encode(
        "utf-8"
    )

    return hashlib.sha256(
        encoded
    ).hexdigest()


def media_plan_path(
    excel_path: Path,
) -> Path:
    return excel_path.with_name(
        f"{excel_path.stem}"
        ".import-plan.json"
    )


def build_media_plan(
    excel_path: Path,
    approved: list[dict],
    results: list[dict],
) -> dict:
    row_by_number = {
        row[
            "_row_number"
        ]: row
        for row in approved
    }

    items = []

    for result in results:
        if result[
            "status"
        ] != "READY":
            continue

        row = row_by_number[
            result[
                "row"
            ]
        ]

        shortcode = clean_cell(
            row,
            "Shortcode",
        )

        if not shortcode:
            raise RuntimeError(
                "Approved row "
                f"{result['row']} "
                "has no Shortcode."
            )

        plan_item = {
            "row": result["row"],
            "shortcode": shortcode,
            "username": result["username"],
            "sourceUrl": result["source_url"],
            "videoCategory": result["video_category"],
            "contentKind": result["content_kind"],
            "creatorSlug": result["creator_slug"],
        }

        raw_items = raw_media_items(
            excel_path,
            plan_item,
            result["content_kind"],
        )

        if not raw_items:
            raise RuntimeError(
                "No downloaded media manifest found for approved row "
                f"{result['row']} ({shortcode}). Run download_approved.py first."
            )

        if result["content_kind"] == "VIDEO" and (
            len(raw_items) != 1
            or raw_items[0]["mediaType"] != "VIDEO"
        ):
            raise RuntimeError(
                f"Approved row {result['row']} is VIDEO but its raw media "
                "does not contain exactly one video item."
            )

        media_items = []
        planned_media = []
        stem = result["identity"]["stem"]

        for index, raw_item in enumerate(raw_items, start=1):
            item_stem = (
                stem
                if result["content_kind"] == "VIDEO"
                else f"{stem}-{index:02d}"
            )
            media_type = raw_item["mediaType"]
            media_url = (
                f"{item_stem}.mp4"
                if media_type == "VIDEO"
                else f"{item_stem}.webp"
            )
            thumbnail_url = (
                f"{item_stem}-thumbnail.webp"
                if media_type == "VIDEO"
                else media_url
            )
            media_items.append(
                {
                    "mediaType": media_type,
                    "mediaUrl": media_url,
                    "thumbnailUrl": thumbnail_url,
                }
            )
            planned_media.append(
                {
                    "displayOrder": index,
                    "mediaType": media_type,
                    "sourceMedia": str(raw_item["mediaPath"]),
                    "sourceThumbnail": (
                        str(raw_item["thumbnailPath"])
                        if raw_item.get("thumbnailPath")
                        else None
                    ),
                    "mediaUrl": media_url,
                    "thumbnailUrl": thumbnail_url,
                }
            )

        payload = dict(result["payload"])
        payload["mediaItems"] = media_items
        payload["mediaUrl"] = media_items[0]["mediaUrl"]
        payload["thumbnailUrl"] = media_items[0]["thumbnailUrl"]
        plan_item["mediaItems"] = planned_media
        plan_item["payload"] = payload
        items.append(plan_item)

    return {
        "version": 2,
        "generatedAt": datetime.now(
            timezone.utc
        ).isoformat(),
        "excel": str(
            excel_path
        ),
        "workbookFingerprint": (
            approved_rows_fingerprint(
                approved
            )
        ),
        "approvedCount": len(
            approved
        ),
        "readyCount": len(
            items
        ),
        "items": items,
    }


def load_media_plan(
    path: Path,
) -> dict:
    return json.loads(
        path.read_text(
            encoding="utf-8"
        )
    )


def write_media_plan(
    path: Path,
    plan: dict,
):
    temporary = path.with_suffix(
        path.suffix + ".tmp"
    )

    temporary.write_text(
        json.dumps(
            plan,
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )

    temporary.replace(
        path
    )


def public_url_to_path(
    value: str,
) -> Path:
    url_path = str(
        value or ""
    ).strip()

    if not url_path.startswith(
        "/"
    ):
        raise RuntimeError(
            "Public media URL must "
            f"start with '/': {url_path}"
        )

    root = PUBLIC_DIR.resolve()
    candidate = (
        root
        / url_path.lstrip(
            "/"
        )
    ).resolve()

    try:
        candidate.relative_to(
            root
        )
    except ValueError as exc:
        raise RuntimeError(
            "Media target escapes "
            "apps/web/public."
        ) from exc

    return candidate


def legacy_raw_media_paths(
    excel_path: Path,
    item: dict,
) -> tuple[Path, Path]:
    username = str(
        item.get(
            "username",
            "",
        )
        or ""
    ).strip()

    shortcode = str(
        item.get(
            "shortcode",
            "",
        )
        or ""
    ).strip()

    base = (
        excel_path.parent
        / username
    )

    return (
        base
        / "videos"
        / f"{shortcode}.mp4",
        base
        / "thumbnails"
        / f"{shortcode}.jpg",
    )


def raw_media_items(
    excel_path: Path,
    item: dict,
    expected_content_kind: str,
) -> list[dict]:
    username = str(item.get("username", "") or "").strip()
    shortcode = str(item.get("shortcode", "") or "").strip()
    folder = excel_path.parent / username / "media" / shortcode
    manifest_path = folder / "media.json"

    if manifest_path.exists():
        manifest = json.loads(
            manifest_path.read_text(encoding="utf-8")
        )
        if manifest.get("version") != 1:
            raise RuntimeError(
                f"Unsupported raw media manifest: {manifest_path}"
            )
        if manifest.get("contentKind") != expected_content_kind:
            raise RuntimeError(
                f"Content kind changed for {shortcode}. Re-run "
                "download_approved.py so its media manifest matches Excel."
            )

        result = []
        for index, media in enumerate(manifest.get("items", []), start=1):
            if not isinstance(media, dict):
                raise RuntimeError(f"Invalid media item in {manifest_path}")
            media_type = str(media.get("mediaType", "")).strip().upper()
            media_path = safe_raw_child_path(
                folder,
                media.get("mediaPath"),
            )
            thumbnail_name = media.get("thumbnailPath")
            thumbnail_path = (
                safe_raw_child_path(folder, thumbnail_name)
                if thumbnail_name
                else None
            )
            if media_type not in {"IMAGE", "VIDEO"}:
                raise RuntimeError(f"Invalid media type in {manifest_path}")
            result.append(
                {
                    "displayOrder": index,
                    "mediaType": media_type,
                    "mediaPath": media_path,
                    "thumbnailPath": thumbnail_path,
                }
            )
        return result

    video_path, thumb_path = legacy_raw_media_paths(excel_path, item)
    if expected_content_kind == "VIDEO" and (
        video_path.exists() or thumb_path.exists()
    ):
        return [
            {
                "displayOrder": 1,
                "mediaType": "VIDEO",
                "mediaPath": video_path,
                "thumbnailPath": thumb_path,
            }
        ]

    return []


def safe_raw_child_path(
    folder: Path,
    value,
) -> Path:
    root = folder.resolve()
    candidate = (
        folder
        / str(value or "")
    ).resolve()

    try:
        candidate.relative_to(root)
    except ValueError as exc:
        raise RuntimeError(
            "Raw media manifest path escapes "
            f"its folder: {value}"
        ) from exc

    return candidate


def validate_raw_media(
    excel_path: Path,
    plan: dict,
):
    missing = []

    for item in plan[
        "items"
    ]:
        if plan.get("version") == 1:
            video_path, thumb_path = legacy_raw_media_paths(excel_path, item)
            media_items = [
                {
                    "mediaType": "VIDEO",
                    "mediaPath": video_path,
                    "thumbnailPath": thumb_path,
                }
            ]
        else:
            media_items = item.get("mediaItems", [])

        for media in media_items:
            media_path = Path(media.get("sourceMedia") or media.get("mediaPath", ""))
            thumbnail_value = media.get("sourceThumbnail") or media.get("thumbnailPath")
            thumbnail_path = Path(thumbnail_value) if thumbnail_value else None
            if not media_path.exists() or media_path.stat().st_size <= 0:
                missing.append(f"Missing media: {media_path}")
            if media.get("mediaType") == "VIDEO" and thumbnail_path and (
                not thumbnail_path.exists() or thumbnail_path.stat().st_size <= 0
            ):
                missing.append(f"Missing thumbnail: {thumbnail_path}")

    if missing:
        raise RuntimeError(
            "Raw media preflight failed:\n"
            + "\n".join(
                f" - {item}"
                for item in missing
            )
        )


def copy_video_file(
    source: Path,
    target: Path,
) -> str:
    target.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if target.exists():
        if (
            target.stat().st_size
            == source.stat().st_size
            and target.stat().st_size
            > 0
        ):
            return "EXISTS"

        raise RuntimeError(
            "Target video already exists "
            "with a different size: "
            f"{target}"
        )

    temporary = target.with_name(
        target.name + ".part"
    )

    if temporary.exists():
        temporary.unlink()

    shutil.copy2(
        source,
        temporary,
    )

    temporary.replace(
        target
    )

    return "COPIED"


def convert_thumbnail_to_webp(
    source: Path,
    target: Path,
) -> str:
    target.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    if (
        target.exists()
        and target.stat().st_size
        > 0
    ):
        return "EXISTS"

    temporary = target.with_name(
        target.stem
        + ".part.webp"
    )

    if temporary.exists():
        temporary.unlink()

    ffmpeg = shutil.which(
        "ffmpeg"
    )

    converted = False

    if ffmpeg:
        process = subprocess.run(
            [
                ffmpeg,
                "-y",
                "-loglevel",
                "error",
                "-i",
                str(source),
                "-frames:v",
                "1",
                "-c:v",
                "libwebp",
                "-quality",
                "90",
                str(temporary),
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.PIPE,
            text=True,
        )

        converted = (
            process.returncode == 0
            and temporary.exists()
            and temporary.stat().st_size
            > 0
        )

        if (
            not converted
            and temporary.exists()
        ):
            temporary.unlink()

    if not converted:
        try:
            from PIL import Image
        except ImportError as exc:
            detail = (
                process.stderr.strip()
                if ffmpeg
                and "process" in locals()
                else ""
            )

            raise RuntimeError(
                "Thumbnail conversion needs "
                "working WebP support in "
                "ffmpeg or Pillow. "
                f"ffmpeg detail: {detail}"
            ) from exc

        with Image.open(
            source
        ) as image:
            image.save(
                temporary,
                format="WEBP",
                quality=90,
            )

    if (
        not temporary.exists()
        or temporary.stat().st_size
        <= 0
    ):
        raise RuntimeError(
            "Thumbnail conversion created "
            "an empty file: "
            f"{temporary}"
        )

    temporary.replace(
        target
    )

    return "CONVERTED"


def verify_prepared_media(
    plan: dict,
) -> tuple[int, int, int, int]:
    videos = 0
    images = 0
    thumbnails = 0
    complete = 0

    for item in plan[
        "items"
    ]:
        for media in planned_media_items(plan, item):
            media_path = public_url_to_path(media["mediaUrl"])
            thumb_path = public_url_to_path(media["thumbnailUrl"])
            has_media = media_path.exists() and media_path.stat().st_size > 0
            has_thumb = thumb_path.exists() and thumb_path.stat().st_size > 0

            if has_media and media["mediaType"] == "VIDEO":
                videos += 1
            if has_media and media["mediaType"] == "IMAGE":
                images += 1
            if has_thumb:
                thumbnails += 1
            if has_media and has_thumb:
                complete += 1

    return (
        videos,
        images,
        thumbnails,
        complete,
    )


def planned_media_items(
    plan: dict,
    item: dict,
) -> list[dict]:
    if plan.get("version") == 2:
        return item.get("mediaItems", [])

    payload = item["payload"]
    return [
        {
            "displayOrder": 1,
            "mediaType": "VIDEO",
            "mediaUrl": payload["mediaUrl"],
            "thumbnailUrl": payload["thumbnailUrl"],
        }
    ]



def find_video_by_id(
    videos: list[dict],
    video_id: str,
):
    target = str(
        video_id or ""
    ).strip()

    for video in videos:
        existing_id = str(
            video.get(
                "id",
                "",
            )
            or ""
        ).strip()

        if (
            existing_id
            and existing_id == target
        ):
            return video

    return None


def validate_plan_structure(
    plan: dict,
):
    if plan.get(
        "version"
    ) not in {1, 2}:
        raise RuntimeError(
            "Unsupported import plan "
            f"version: {plan.get('version')}"
        )

    items = plan.get(
        "items"
    )

    if not isinstance(
        items,
        list,
    ):
        raise RuntimeError(
            "Import plan has no valid "
            "items list."
        )

    if (
        plan.get(
            "readyCount"
        )
        != len(items)
    ):
        raise RuntimeError(
            "Import plan readyCount does "
            "not match its item count."
        )

    seen_ids = set()
    seen_sources = set()
    seen_media = set()
    seen_thumbnails = set()

    for index, item in enumerate(
        items,
        start=1,
    ):
        payload = item.get(
            "payload"
        )

        if not isinstance(
            payload,
            dict,
        ):
            raise RuntimeError(
                "Import plan item "
                f"{index} has no payload."
            )

        video_id = str(
            payload.get(
                "id",
                "",
            )
            or ""
        ).strip()

        source_url = str(
            payload.get(
                "sourceUrl",
                "",
            )
            or ""
        ).strip().rstrip(
            "/"
        )

        media_url = str(payload.get("mediaUrl", "") or "").strip()
        thumbnail_url = str(payload.get("thumbnailUrl", "") or "").strip()

        required = {
            "id": video_id,
            "sourceUrl": source_url,
            "mediaUrl": media_url,
            "thumbnailUrl": thumbnail_url,
        }

        for name, value in required.items():
            if not value:
                raise RuntimeError(
                    "Import plan item "
                    f"{index} has empty "
                    f"{name}."
                )

        duplicate_sets = [
            (
                "video id",
                video_id,
                seen_ids,
            ),
            (
                "source URL",
                source_url,
                seen_sources,
            ),
        ]

        if plan.get("version") == 1:
            duplicate_sets.extend(
                [
                    ("media URL", media_url, seen_media),
                    ("thumbnail URL", thumbnail_url, seen_thumbnails),
                ]
            )

        for (
            label_name,
            value,
            seen,
        ) in duplicate_sets:
            if value in seen:
                raise RuntimeError(
                    "Import plan contains "
                    f"duplicate {label_name}: "
                    f"{value}"
                )

            seen.add(
                value
            )

        if plan.get("version") == 1:
            continue

        planned = planned_media_items(plan, item)
        if not planned:
            raise RuntimeError(
                f"Import plan item {index} has no media items."
            )
        content_kind = payload.get("contentKind")
        if content_kind not in {"VIDEO", "POST", "STORY"}:
            raise RuntimeError(
                f"Import plan item {index} has invalid contentKind."
            )
        if content_kind == "VIDEO" and (
            len(planned) != 1
            or planned[0].get("mediaType") != "VIDEO"
        ):
            raise RuntimeError(
                f"Import plan item {index} violates VIDEO media rules."
            )

        for media_index, media in enumerate(planned, start=1):
            media_type = media.get("mediaType")
            item_media_url = str(media.get("mediaUrl", "") or "").strip()
            item_thumbnail_url = str(
                media.get("thumbnailUrl", "") or ""
            ).strip()
            if media_type not in {"IMAGE", "VIDEO"}:
                raise RuntimeError(
                    f"Import plan item {index}.{media_index} has invalid mediaType."
                )
            if not item_media_url or not item_thumbnail_url:
                raise RuntimeError(
                    f"Import plan item {index}.{media_index} has empty paths."
                )
            for label_name, value, seen in (
                ("media URL", item_media_url, seen_media),
                ("thumbnail URL", item_thumbnail_url, seen_thumbnails),
            ):
                if value in seen and not (
                    media_type == "IMAGE"
                    and label_name == "thumbnail URL"
                    and value == item_media_url
                ):
                    raise RuntimeError(
                        f"Import plan contains duplicate {label_name}: {value}"
                    )
                seen.add(value)


def validate_plan_workbook(
    excel_path: Path,
    plan: dict,
):
    rows = load_excel_rows(
        excel_path
    )

    approved = [
        row
        for row in rows
        if is_approved(row)
    ]

    fingerprint = (
        approved_rows_fingerprint(
            approved
        )
    )

    if (
        plan.get(
            "workbookFingerprint"
        )
        != fingerprint
    ):
        raise RuntimeError(
            "The reviewed workbook changed "
            "after the stable import plan "
            "was created. Run "
            "--prepare-media again only "
            "after intentionally rebuilding "
            "the plan."
        )

    if (
        plan.get(
            "approvedCount"
        )
        != len(approved)
    ):
        raise RuntimeError(
            "Import plan approvedCount does "
            "not match the workbook."
        )


def validate_prepared_media_strict(
    plan: dict,
):
    missing = []

    for item in plan[
        "items"
    ]:
        for media in planned_media_items(plan, item):
            media_path = public_url_to_path(media["mediaUrl"])
            thumb_path = public_url_to_path(media["thumbnailUrl"])
            if not media_path.exists() or media_path.stat().st_size <= 0:
                missing.append(f"Missing final media: {media_path}")
            if not thumb_path.exists() or thumb_path.stat().st_size <= 0:
                missing.append(f"Missing final thumbnail: {thumb_path}")

    if missing:
        raise RuntimeError(
            "Prepared media preflight failed:\n"
            + "\n".join(
                f" - {item}"
                for item in missing
            )
        )


def validate_plan_catalog_refs(
    plan: dict,
    catalog: dict,
):
    people_handles = {
        normalize_handle(
            person.get(
                "instagramHandle"
            )
            or person.get(
                "instagram_handle"
            )
            or ""
        )
        for person in catalog[
            "people"
        ]
    }

    destination_ids = {
        str(
            item.get(
                "id",
                "",
            )
            or ""
        ).strip()
        for item in catalog[
            "destinations"
        ]
    }

    hotel_ids = {
        str(
            item.get(
                "id",
                "",
            )
            or ""
        ).strip()
        for item in catalog[
            "hotels"
        ]
    }

    problems = []

    for item in plan[
        "items"
    ]:
        payload = item[
            "payload"
        ]

        username = normalize_handle(
            payload.get(
                "instagramUsername"
            )
        )

        if (
            not username
            or username
            not in people_handles
        ):
            problems.append(
                "Missing person for "
                f"@{payload.get('instagramUsername', '')}"
            )

        for destination_id in payload.get(
            "destinationIds",
            [],
        ):
            if (
                str(destination_id)
                not in destination_ids
            ):
                problems.append(
                    "Missing destination id "
                    f"{destination_id} for "
                    f"{payload['id']}"
                )

        for hotel_id in payload.get(
            "hotelIds",
            [],
        ):
            if (
                str(hotel_id)
                not in hotel_ids
            ):
                problems.append(
                    "Missing hotel id "
                    f"{hotel_id} for "
                    f"{payload['id']}"
                )

    if problems:
        raise RuntimeError(
            "Catalog reference preflight failed:\n"
            + "\n".join(
                f" - {item}"
                for item in problems
            )
        )


def classify_plan_against_catalog(
    plan: dict,
    catalog: dict,
) -> tuple[list[dict], list[dict]]:
    to_create = []
    existing = []

    for item in plan[
        "items"
    ]:
        payload = item[
            "payload"
        ]

        source_match = find_video_by_source(
            catalog[
                "videos"
            ],
            payload[
                "sourceUrl"
            ],
        )

        if source_match:
            existing.append(
                item
            )
            continue

        id_match = find_video_by_id(
            catalog[
                "videos"
            ],
            payload[
                "id"
            ],
        )

        if id_match:
            existing_source = str(
                id_match.get(
                    "sourceUrl",
                    "",
                )
                or ""
            ).strip()

            raise RuntimeError(
                "Video ID collision before "
                "apply: "
                f"{payload['id']} already "
                "exists with source "
                f"{existing_source or 'unknown'}."
            )

        to_create.append(
            item
        )

    return (
        to_create,
        existing,
    )


def run_apply(
    excel_path: Path,
):
    plan_path = media_plan_path(
        excel_path
    )

    if not plan_path.exists():
        raise RuntimeError(
            "Stable import plan not found. "
            "Run --prepare-media first."
        )

    print(
        "Loading stable import plan..."
    )

    plan = load_media_plan(
        plan_path
    )

    validate_plan_structure(
        plan
    )

    validate_plan_workbook(
        excel_path,
        plan,
    )

    validate_prepared_media_strict(
        plan
    )

    print(
        "Loading Hotel-Yab catalog..."
    )

    catalog = load_catalog()

    validate_plan_catalog_refs(
        plan,
        catalog,
    )

    to_create, existing = (
        classify_plan_against_catalog(
            plan,
            catalog,
        )
    )

    print()
    print(
        color(
            (
                "════════════════ "
                "APPLY PREFLIGHT "
                "════════════════"
            ),
            Color.BOLD
            + Color.WHITE,
        )
    )
    print()

    label(
        "Plan items",
        str(
            len(
                plan[
                    "items"
                ]
            )
        ),
    )
    label(
        "To create",
        color(
            str(
                len(to_create)
            ),
            Color.GREEN,
        ),
    )
    label(
        "Existing",
        color(
            str(
                len(existing)
            ),
            Color.CYAN,
        ),
    )
    label(
        "Media pairs",
        color(
            str(
                len(
                    plan[
                        "items"
                    ]
                )
            ),
            Color.GREEN,
        ),
    )

    if not to_create:
        print()
        print(
            color(
                "Nothing new to create.",
                Color.CYAN,
            )
        )
        return

    print()
    print(
        color(
            (
                "Creating videos from the "
                "stable import plan..."
            ),
            Color.WHITE
            + Color.BOLD,
        )
    )

    created = 0
    failed = 0

    for index, item in enumerate(
        to_create,
        start=1,
    ):
        payload = item[
            "payload"
        ]

        print(
            f"[{index}/{len(to_create)}] "
            f"{payload['id']}"
        )

        try:
            api_post_json(
                "/admin/catalog/videos",
                payload,
            )
        except RuntimeError as exc:
            failed += 1

            print(
                color(
                    (
                        "   ❌ FAILED: "
                        f"{exc}"
                    ),
                    Color.RED,
                )
            )

            print()
            print(
                color(
                    (
                        "Apply stopped after the "
                        "first failure. Re-run "
                        "--apply after fixing the "
                        "problem; already-created "
                        "sources will be skipped."
                    ),
                    Color.YELLOW,
                )
            )
            break

        created += 1

        print(
            color(
                "   ✅ CREATED",
                Color.GREEN,
            )
        )

    print()
    print(
        color(
            (
                "════════════════ "
                "APPLY SUMMARY "
                "════════════════"
            ),
            Color.BOLD
            + Color.WHITE,
        )
    )
    print()

    label(
        "Plan items",
        str(
            len(
                plan[
                    "items"
                ]
            )
        ),
    )
    label(
        "Created",
        color(
            str(created),
            Color.GREEN
            + Color.BOLD,
        ),
    )
    label(
        "⏭ Existing",
        color(
            str(
                len(existing)
            ),
            Color.CYAN,
        ),
    )
    label(
        "❌ Failed",
        color(
            str(failed),
            (
                Color.RED
                if failed
                else Color.GREEN
            ),
        ),
    )

    if failed:
        raise RuntimeError(
            "Apply did not complete. "
            "No automatic retry was "
            "attempted."
        )

    print()
    print(
        color(
            (
                "All planned new videos "
                "were created successfully."
            ),
            Color.GREEN
            + Color.BOLD,
        )
    )


def run_prepare_media(
    excel_path: Path,
):
    print(
        "Loading Hotel-Yab catalog..."
    )

    catalog = load_catalog()

    approved, results = (
        analyze_approved_rows(
            excel_path,
            catalog,
        )
    )

    blocked = [
        result
        for result in results
        if result[
            "status"
        ] == "BLOCKED"
    ]

    if blocked:
        for result in blocked:
            print_result(
                result
            )

        raise RuntimeError(
            "Media preparation stopped "
            "because approved rows are "
            "blocked. Run --dry-run first."
        )

    current_fingerprint = (
        approved_rows_fingerprint(
            approved
        )
    )

    plan_path = media_plan_path(
        excel_path
    )

    if plan_path.exists():
        plan = load_media_plan(
            plan_path
        )

        if (
            plan.get(
                "workbookFingerprint"
            )
            != current_fingerprint
        ):
            raise RuntimeError(
                "The reviewed workbook changed "
                "after the media plan was created. "
                f"Delete {plan_path.name} and run "
                "--prepare-media again to create "
                "a fresh plan."
            )

        print(
            "Using existing stable "
            f"media plan: {plan_path}"
        )

    else:
        plan = build_media_plan(
            excel_path,
            approved,
            results,
        )

        validate_raw_media(
            excel_path,
            plan,
        )

        write_media_plan(
            plan_path,
            plan,
        )

        print(
            "Stable media plan saved:"
        )
        print(
            plan_path
        )

    validate_raw_media(
        excel_path,
        plan,
    )

    copied_videos = 0
    existing_videos = 0
    converted_images = 0
    existing_images = 0
    converted_thumbnails = 0
    existing_thumbnails = 0

    for index, item in enumerate(
        plan[
            "items"
        ],
        start=1,
    ):
        payload = item[
            "payload"
        ]

        print(
            f"[{index}/{len(plan['items'])}] "
            f"{item['shortcode']} "
            f"→ {payload['id']}"
        )

        if plan.get("version") == 1:
            source_video, source_thumb = legacy_raw_media_paths(
                excel_path,
                item,
            )
            media_entries = [
                {
                    "displayOrder": 1,
                    "mediaType": "VIDEO",
                    "sourceMedia": str(source_video),
                    "sourceThumbnail": str(source_thumb),
                    "mediaUrl": payload["mediaUrl"],
                    "thumbnailUrl": payload["thumbnailUrl"],
                }
            ]
        else:
            media_entries = item["mediaItems"]

        for media in media_entries:
            source_media = Path(media["sourceMedia"])
            source_thumbnail = (
                Path(media["sourceThumbnail"])
                if media.get("sourceThumbnail")
                else source_media
            )
            target_media = public_url_to_path(media["mediaUrl"])
            target_thumbnail = public_url_to_path(media["thumbnailUrl"])

            if media["mediaType"] == "VIDEO":
                media_status = copy_video_file(source_media, target_media)
                thumbnail_status = convert_thumbnail_to_webp(
                    source_thumbnail,
                    target_thumbnail,
                )
                if media_status == "COPIED":
                    copied_videos += 1
                else:
                    existing_videos += 1
                if thumbnail_status == "CONVERTED":
                    converted_thumbnails += 1
                else:
                    existing_thumbnails += 1
            else:
                media_status = convert_thumbnail_to_webp(
                    source_media,
                    target_media,
                )
                thumbnail_status = media_status
                if media_status == "CONVERTED":
                    converted_images += 1
                else:
                    existing_images += 1

            print(
                f"   {media['displayOrder']:02d} {media['mediaType']}: "
                f"{media['mediaUrl']} ({media_status})"
            )
            if media["mediaType"] == "VIDEO":
                print(
                    f"      Thumb: {media['thumbnailUrl']} "
                    f"({thumbnail_status})"
                )

    final_videos, final_images, final_thumbs, complete = (
        verify_prepared_media(
            plan
        )
    )

    print()
    print(
        color(
            (
                "════════════════ "
                "MEDIA PREP SUMMARY "
                "════════════════"
            ),
            Color.BOLD
            + Color.WHITE,
        )
    )
    print()

    label(
        "Approved",
        str(
            len(approved)
        ),
    )
    label(
        "Plan items",
        str(
            len(
                plan[
                    "items"
                ]
            )
        ),
    )
    label(
        "Videos copied",
        str(
            copied_videos
        ),
    )
    label(
        "Videos existing",
        str(
            existing_videos
        ),
    )
    label(
        "Images made",
        str(converted_images),
    )
    label(
        "Images existing",
        str(existing_images),
    )
    label(
        "Thumbs made",
        str(
            converted_thumbnails
        ),
    )
    label(
        "Thumbs existing",
        str(
            existing_thumbnails
        ),
    )
    label(
        "Final videos",
        str(
            final_videos
        ),
    )
    label(
        "Final thumbs",
        str(
            final_thumbs
        ),
    )
    label(
        "Final images",
        str(final_images),
    )
    expected_media = sum(
        len(planned_media_items(plan, item))
        for item in plan["items"]
    )
    label(
        "Complete media",
        color(
            str(
                complete
            ),
            (
                Color.GREEN
                if complete
                == expected_media
                else Color.RED
            ),
        ),
    )

    print()
    print(
        "Plan manifest:"
    )
    print(
        plan_path
    )
    print()
    print(
        color(
            (
                "No database changes "
                "were made."
            ),
            Color.GRAY,
        )
    )

def run_dry_run(
    excel_path: Path,
):
    global catalog_for_print
    print(
        "Loading Hotel-Yab catalog..."
    )

    catalog = load_catalog()
    catalog_for_print = catalog

    print(
        "Catalog:"
    )
    print(
        " Destinations:",
        len(
            catalog[
                "destinations"
            ]
        ),
    )
    print(
        " Hotels:",
        len(
            catalog["hotels"]
        ),
    )
    print(
        " People:",
        len(
            catalog["people"]
        ),
    )
    print(
        " Videos:",
        len(
            catalog["videos"]
        ),
    )

    rows = load_excel_rows(
        excel_path
    )

    approved = [
        row
        for row in rows
        if is_approved(row)
    ]

    print(
        "\nApproved rows:",
        len(approved),
    )

    sequence_state = {}
    batch_sources = set()
    used_titles = existing_video_title_keys(catalog["videos"])

    results = [
        analyze_row(
            row,
            catalog,
            sequence_state,
            batch_sources,
            used_titles,
        )
        for row in approved
    ]

    for result in results:
        print_result(result)

    ready = sum(
        result["status"]
        == "READY"
        for result in results
    )

    blocked = sum(
        result["status"]
        == "BLOCKED"
        for result in results
    )

    existing = sum(
        result["status"]
        == "ALREADY_EXISTS"
        for result in results
    )

    print()

    print(
        color(
            (
                "════════════════ "
                "DRY RUN SUMMARY "
                "════════════════"
            ),
            Color.BOLD
            + Color.WHITE,
        )
    )

    print()

    label(
        "Approved",
        color(
            str(
                len(results)
            ),
            Color.WHITE,
        ),
    )

    label(
        "✅ Ready",
        color(
            str(ready),
            Color.GREEN
            + Color.BOLD,
        ),
    )

    label(
        "❌ Blocked",
        color(
            str(blocked),
            Color.RED
            + Color.BOLD,
        ),
    )

    label(
        "⏭ Existing",
        color(
            str(existing),
            Color.CYAN,
        ),
    )

    print()

    if blocked:
        print(
            color(
                (
                    f"{blocked} row(s) "
                    "need attention "
                    "before --apply."
                ),
                Color.RED,
            )
        )

    elif ready:
        print(
            color(
                (
                    "All new approved "
                    "rows are ready."
                ),
                Color.GREEN,
            )
        )

    print()

    print(
        color(
            (
                "No database changes "
                "were made."
            ),
            Color.GRAY,
        )
    )



def main():
    parser = (
        argparse.ArgumentParser(
            description=(
                "Validate approved "
                "Instagram video rows, "
                "prepare their media, "
                "or apply the stable plan "
                "to Hotel-Yab."
            )
        )
    )

    parser.add_argument(
        "excel",
        help=(
            "Reviewed XLSX/JSON path or "
            "profile name such as "
            "minaaaslife"
        ),
    )

    mode = parser.add_mutually_exclusive_group(
        required=True
    )

    mode.add_argument(
        "--dry-run",
        action="store_true",
        help=(
            "Validate without "
            "creating anything"
        ),
    )

    mode.add_argument(
        "--prepare-media",
        action="store_true",
        help=(
            "Copy approved MP4 files "
            "to their final public paths "
            "and convert thumbnails "
            "to WebP"
        ),
    )

    mode.add_argument(
        "--apply",
        action="store_true",
        help=(
            "Create videos through the "
            "Hotel-Yab Admin API using "
            "the stable import plan"
        ),
    )

    args = parser.parse_args()

    excel_path = resolve_excel_path(
        args.excel
    )

    if args.dry_run:
        run_dry_run(
            excel_path
        )
        return

    if args.prepare_media:
        run_prepare_media(
            excel_path
        )
        return

    if args.apply:
        run_apply(
            excel_path
        )
        return


if __name__ == "__main__":
    main()
